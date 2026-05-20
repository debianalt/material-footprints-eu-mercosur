"""
md_to_docx.py  — conversión manuscript_draft.md → .docx (estilo Ecological Economics)
- Times New Roman 12pt, doble espaciado, justificado
- Line numbers continuos
- Fórmulas: subíndices/superíndices reales, centradas, itálica
- Color: negro en todos los runs (sin colores)
"""

import re
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# BASE = carpeta del manuscrito, derivada de la ubicación de este script
# (tools/ está dentro de BASE). Evita romperse si la carpeta se renombra.
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC  = os.path.join(BASE, "manuscript_draft.md")
DST  = os.path.join(BASE, "output", "manuscript_EE_submission.docx")

BLACK = RGBColor(0, 0, 0)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def set_run_base(run, size=12, bold=False, italic=False, subscript=False, superscript=False):
    run.font.name      = "Times New Roman"
    run.font.size      = Pt(size)
    run.font.bold      = bold
    run.font.italic    = italic
    run.font.color.rgb = BLACK
    if subscript:
        run.font.subscript   = True
    if superscript:
        run.font.superscript = True

def set_para_normal(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    para.alignment = align
    fmt = para.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_after  = Pt(0)
    fmt.space_before = Pt(0)


# ---------------------------------------------------------------------------
# Formula run builder — handles _{...} ^{...} _x ^x inline subscripts
# ---------------------------------------------------------------------------
FORMULA_TOKEN = re.compile(
    r"_\{([^}]+)\}"        # _{subscript}
    r"|\^\{([^}]+)\}"      # ^{superscript}
    r"|_([A-Za-z0-9]+)"    # _word multi-char subscript, no braces (gap_rel, g_GDP)
    r"|\^([A-Za-z0-9]+)"   # ^word multi-char superscript
    r"|([^_^{}]+)"         # plain text chunk
)

def add_formula_runs(para, text, body_size=12, formula_mode=False):
    """
    Add runs to `para` parsing inline subscripts/superscripts.
    formula_mode=True → all runs italic (for display equations).
    """
    for m in FORMULA_TOKEN.finditer(text):
        sub_brace  = m.group(1)  # _{...}
        sup_brace  = m.group(2)  # ^{...}
        sub_single = m.group(3)  # _x
        sup_single = m.group(4)  # ^x
        plain      = m.group(5)  # normal text

        if sub_brace is not None:
            r = para.add_run(sub_brace)
            set_run_base(r, size=body_size - 2, italic=True, subscript=True)
        elif sup_brace is not None:
            r = para.add_run(sup_brace)
            set_run_base(r, size=body_size - 2, italic=True, superscript=True)
        elif sub_single is not None:
            r = para.add_run(sub_single)
            set_run_base(r, size=body_size - 2, italic=True, subscript=True)
        elif sup_single is not None:
            r = para.add_run(sup_single)
            set_run_base(r, size=body_size - 2, italic=True, superscript=True)
        else:
            r = para.add_run(plain)
            set_run_base(r, size=body_size, italic=formula_mode)


# ---------------------------------------------------------------------------
# Inline markdown formatter — **bold**, *italic*, `code`, + formula tokens
# ---------------------------------------------------------------------------
INLINE_TOKEN = re.compile(
    r"\*\*(.+?)\*\*"   # **bold**
    r"|\*(.+?)\*"      # *italic*
    r"|`(.+?)`"        # `code`
)

def add_inline_runs(para, text, size=12):
    """
    Parse **bold**, *italic*, `code` and formula subscripts in a paragraph run.
    For each plain-text segment between markdown marks, further parse formula tokens.
    """
    last = 0
    for m in INLINE_TOKEN.finditer(text):
        if m.start() > last:
            add_formula_runs(para, text[last:m.start()], body_size=size)
        bold_txt  = m.group(1)
        ital_txt  = m.group(2)
        code_txt  = m.group(3)
        if bold_txt is not None:
            r = para.add_run(bold_txt)
            set_run_base(r, size=size, bold=True)
        elif ital_txt is not None:
            r = para.add_run(ital_txt)
            set_run_base(r, size=size, italic=True)
        else:
            r = para.add_run(code_txt)
            r.font.name      = "Courier New"
            r.font.size      = Pt(size - 1)
            r.font.color.rgb = BLACK
        last = m.end()
    if last < len(text):
        add_formula_runs(para, text[last:], body_size=size)


# ---------------------------------------------------------------------------
# Table builder
# ---------------------------------------------------------------------------
def is_table_line(line):
    return line.strip().startswith("|")

def is_sep_line(line):
    return bool(re.match(r"^\|[\s\-\|:]+\|$", line.strip()))

def parse_table(lines):
    # Split each row on unescaped `|`; treat `\|` as a literal pipe inside a cell
    # (needed for math notation such as `\|ε − 1\| ≤ τ`).
    rows = []
    for ln in lines:
        if is_sep_line(ln):
            continue
        body = ln.strip().strip("|")
        cells = [c.replace(r"\|", "|").strip() for c in re.split(r"(?<!\\)\|", body)]
        rows.append(cells)
    return rows

def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows  = [r + [""] * (ncols - len(r)) for r in rows]
    tbl   = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = tbl.rows[i].cells[j]
            p    = cell.paragraphs[0]
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after  = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            # header row bold, otherwise normal
            r = p.add_run(txt)
            r.font.name      = "Times New Roman"
            r.font.size      = Pt(10)
            r.font.bold      = (i == 0)
            r.font.color.rgb = BLACK
    # small gap after table
    gap = doc.add_paragraph()
    set_para_normal(gap)
    gap.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------
# Detect display equation line
# A standalone line is a display equation if it contains _{, ^{, or _letter
# and looks like math (short, no long English phrases).
# ---------------------------------------------------------------------------
def is_display_equation(line):
    stripped = line.strip()
    if not stripped:
        return False
    # Bold-led lines are labels/captions/notes (e.g. **Figure 6.** ..., **Table 2. ...**,
    # **Specification.**), never display equations.
    if stripped.startswith("**"):
        return False
    has_math = bool(re.search(r"_\{|_[A-Za-z]|\^\{|\^[A-Za-z]", stripped))
    # Must not look like a normal sentence (no article/preposition words)
    looks_english = bool(re.search(r"\b(the|is|are|and|in|of|to|for|we|this|that|which)\b", stripped))
    return has_math and not looks_english


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------
doc = Document()

sec = doc.sections[0]
for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, attr, Cm(2.5))

# Continuous line numbers
def add_line_numbers(section):
    sp = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sp.append(ln)

add_line_numbers(sec)

# Normal style
normal = doc.styles["Normal"]
normal.font.name       = "Times New Roman"
normal.font.size       = Pt(12)
normal.font.color.rgb  = BLACK
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
normal.paragraph_format.space_after       = Pt(0)
normal.paragraph_format.space_before      = Pt(0)
normal.paragraph_format.alignment         = WD_ALIGN_PARAGRAPH.JUSTIFY

# Heading styles
for lvl, sz in [(1, 14), (2, 13), (3, 12)]:
    h = doc.styles[f"Heading {lvl}"]
    h.font.name       = "Times New Roman"
    h.font.size       = Pt(sz)
    h.font.bold       = True
    h.font.color.rgb  = BLACK
    h.paragraph_format.alignment         = WD_ALIGN_PARAGRAPH.LEFT
    h.paragraph_format.space_before      = Pt(12)
    h.paragraph_format.space_after       = Pt(6)
    h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

# ---------------------------------------------------------------------------
# Read source
# ---------------------------------------------------------------------------
with open(SRC, encoding="utf-8") as f:
    src_lines = f.read().splitlines()

# ---------------------------------------------------------------------------
# Parse loop
# ---------------------------------------------------------------------------
i = 0
in_table   = False
table_buf  = []

while i < len(src_lines):
    line = src_lines[i]

    # --- Blank line ---
    if line.strip() == "":
        if in_table:
            add_table(doc, parse_table(table_buf))
            table_buf = []
            in_table  = False
        i += 1
        continue

    # --- Author block IS kept in the manuscript ---
    #     Ecological Economics uses SINGLE-anonymized review: per Elsevier,
    #     the title page (authors, affiliations, corresponding author, ORCID)
    #     is the first page of the manuscript file, not removed from it.
    #     (A separate title-page file is the DOUBLE-anonymized requirement.)

    # --- Horizontal rule ---
    if re.match(r"^-{3,}$", line.strip()):
        if in_table:
            add_table(doc, parse_table(table_buf))
            table_buf = []
            in_table  = False
        p = doc.add_paragraph()
        set_para_normal(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run("─" * 50)
        set_run_base(r, size=10)
        i += 1
        continue

    # --- Image embed: insert the figure inline (integrated submission).
    #     Figures are ALSO supplied as separate TIFFs; the embed gives the
    #     editor a self-contained manuscript. PNG resolved relative to BASE. ---
    mimg = re.match(r"^!\[[^\]]*\]\(([^)]*)\)\s*$", line.strip())
    if mimg:
        img_rel = mimg.group(1).strip()
        img_path = img_rel if os.path.isabs(img_rel) else os.path.join(BASE, img_rel)
        if os.path.isfile(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run()
            try:
                run.add_picture(img_path, width=Cm(15.0))
            except Exception as exc:                       # noqa: BLE001
                p.add_run(f"[figure not embedded: {os.path.basename(img_path)} — {exc}]")
        else:
            p = doc.add_paragraph()
            set_para_normal(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_run_base(p.add_run(f"[missing figure file: {img_rel}]"), size=10, italic=True)
        i += 1
        continue

    # --- Table line ---
    if is_table_line(line):
        in_table = True
        table_buf.append(line)
        i += 1
        continue
    elif in_table:
        add_table(doc, parse_table(table_buf))
        table_buf = []
        in_table  = False

    # --- Heading 1 ---
    m = re.match(r"^# (.+)$", line)
    if m:
        p = doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(m.group(1))
        set_run_base(r, size=14, bold=True)
        i += 1
        continue

    # --- Heading 2 ---
    m = re.match(r"^## (.+)$", line)
    if m:
        p = doc.add_paragraph(style="Heading 2")
        r = p.add_run(m.group(1))
        set_run_base(r, size=13, bold=True)
        i += 1
        continue

    # --- Heading 3 ---
    m = re.match(r"^### (.+)$", line)
    if m:
        p = doc.add_paragraph(style="Heading 3")
        r = p.add_run(m.group(1))
        set_run_base(r, size=12, bold=True)
        i += 1
        continue

    # --- Display equation (standalone math line) ---
    if is_display_equation(line.strip()):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(6)
        add_formula_runs(p, line.strip(), body_size=12, formula_mode=True)
        i += 1
        continue

    # --- Bullet list item ("- text"; "---" already handled above) ---
    m = re.match(r"^- (.+)$", line.strip())
    if m:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_after  = Pt(0)
        pf.space_before = Pt(0)
        pf.left_indent  = Cm(0.75)
        pf.first_line_indent = Cm(-0.75)  # hanging indent
        bullet = p.add_run("•\t")
        set_run_base(bullet, size=12)
        add_inline_runs(p, m.group(1), size=12)
        i += 1
        continue

    # --- Collect paragraph lines ---
    para_lines = []
    while i < len(src_lines):
        ln = src_lines[i]
        if ln.strip() == "":
            break
        if re.match(r"^#{1,3} |^---+$", ln) or is_table_line(ln):
            break
        if is_display_equation(ln.strip()):
            break
        para_lines.append(ln)
        i += 1

    if not para_lines:
        i += 1
        continue

    combined = " ".join(para_lines)

    # Detect block-level *italic note* (starts and ends with single *)
    if (combined.strip().startswith("*") and combined.strip().endswith("*")
            and not combined.strip().startswith("**")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(4)
        inner = combined.strip().strip("*")
        # parse subscripts/superscripts so g_MF, g_GDP, p_self render
        # correctly in table notes (formula_mode keeps the note italic)
        add_formula_runs(p, inner, body_size=10, formula_mode=True)
    else:
        p = doc.add_paragraph()
        set_para_normal(p)
        add_inline_runs(p, combined, size=12)

# Flush trailing table
if in_table:
    add_table(doc, parse_table(table_buf))

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
os.makedirs(os.path.dirname(DST), exist_ok=True)
doc.save(DST)
print(f"Guardado: {DST}")
